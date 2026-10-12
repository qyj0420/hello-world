#!/usr/bin/env python3
"""Generate the privacy-minimized device asset catalog from the factory workbook.

The workbook is opened read-only. JSON is written only to stdout so callers can
review/apply the generated artifact without ever modifying the source workbook.
"""

import argparse
import collections
import datetime as dt
import hashlib
import json
import re
import sys
from pathlib import Path

import openpyxl


PICO_SERIAL = re.compile(r"(?i)(?<![A-Z0-9])([A-Z0-9]{5}/[0-9]{4})(?![A-Z0-9])")
PICO_MODEL = re.compile(r"(?i)PICO(?:SCOPE)?\s*([0-9]{4}[A-Z])")
ART_SERIAL = re.compile(r"(?<![0-9])([0-9]{10})(?![0-9])")
ART_MODEL = re.compile(r"(?i)(ART[0-9]{4}A?)")
NUMBER = re.compile(r"[-+]?[0-9]+(?:\.[0-9]+)?")


def clean(value):
    return "" if value is None else str(value).strip()


def normalize_key(value):
    return re.sub(r"\s+", "", clean(value)).upper().replace("／", "/")


def first_match(pattern, value):
    match = pattern.search(clean(value))
    return normalize_key(match.group(1)) if match else ""


def number(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    match = NUMBER.search(clean(value).replace(",", ""))
    return float(match.group(0)) if match else None


def capacitance_nf(value):
    text = clean(value).replace(" ", "")
    parsed = number(text)
    if parsed is None:
        return None
    lowered = text.lower()
    if "pf" in lowered:
        return parsed / 1000.0
    if "uf" in lowered or "μf" in lowered or "µf" in lowered:
        return parsed * 1000.0
    return parsed


def voltage_kv(value):
    text = clean(value).replace(" ", "")
    parsed = number(text)
    if parsed is None:
        return None
    return parsed / 1000.0 if re.search(r"(?i)(?<!k)v$", text) else parsed


def inductor(value):
    text = clean(value).replace("Ω", "欧姆")
    inductance = re.search(r"(?i)([0-9]+(?:\.[0-9]+)?)\s*(mH|H)", text)
    resistance = re.search(r"([0-9]+(?:\.[0-9]+)?)\s*欧姆", text)
    if not inductance or not resistance:
        return {"inductanceHenry": None, "dcResistanceOhm": None}
    henry = float(inductance.group(1))
    if inductance.group(2).lower() == "mh":
        henry /= 1000.0
    return {"inductanceHenry": henry, "dcResistanceOhm": float(resistance.group(1))}


def boolean_text(value):
    text = clean(value).lower()
    if text in ("是", "有", "yes", "true", "1"):
        return True
    if text in ("否", "无", "no", "false", "0"):
        return False
    return None


def add_flag(flags, condition, name):
    if condition:
        flags.append(name)


def generate(source_path):
    source = Path(source_path)
    workbook = openpyxl.load_workbook(str(source), read_only=True, data_only=True)
    try:
        sheet = workbook["出厂记录"]
        records = []
        for row in range(2, 88):
            factory = normalize_key(sheet.cell(row, 1).value)
            pico_text = clean(sheet.cell(row, 9).value)
            art_text = clean(sheet.cell(row, 10).value)
            compensation_selected = boolean_text(sheet.cell(row, 18).value)
            records.append(
                {
                    "sourceRow": row,
                    "factoryNumber": factory,
                    "deviceModel": clean(sheet.cell(row, 4).value),
                    "picoScope": {
                        "model": first_match(PICO_MODEL, pico_text),
                        "serial": first_match(PICO_SERIAL, pico_text),
                    },
                    "art": {
                        "model": first_match(ART_MODEL, art_text),
                        "serial": first_match(ART_SERIAL, art_text),
                    },
                    "dividerRatio": number(sheet.cell(row, 13).value),
                    "inductor": inductor(sheet.cell(row, 12).value),
                    "couplingCapacitanceNanoFarad": capacitance_nf(sheet.cell(row, 11).value),
                    "compensation": {
                        "selected": compensation_selected,
                        # The source headers are reversed relative to the populated units;
                        # preserve the actual nF/kV semantics rather than their positions.
                        "capacitanceNanoFarad": capacitance_nf(sheet.cell(row, 19).value),
                        "voltageKilovolt": voltage_kv(sheet.cell(row, 20).value),
                    },
                    "calibrationCoefficientSource": clean(sheet.cell(row, 21).value),
                    "calibrationCoefficient": number(sheet.cell(row, 21).value),
                }
            )
    finally:
        workbook.close()

    factory_counts = collections.Counter(r["factoryNumber"] for r in records if r["factoryNumber"])
    pico_counts = collections.Counter(r["picoScope"]["serial"] for r in records if r["picoScope"]["serial"])
    art_counts = collections.Counter(r["art"]["serial"] for r in records if r["art"]["serial"])

    pico_links = collections.defaultdict(set)
    art_links = collections.defaultdict(set)
    for record in records:
        pico = record["picoScope"]["serial"]
        art = record["art"]["serial"]
        if pico and art:
            pico_links[pico].add((record["factoryNumber"], art))
            art_links[art].add((record["factoryNumber"], pico))

    exclusion_counts = collections.Counter()
    for record in records:
        flags = []
        factory = record["factoryNumber"]
        pico = record["picoScope"]["serial"]
        art = record["art"]["serial"]
        add_flag(flags, not factory, "MissingFactoryNumber")
        add_flag(flags, factory_counts.get(factory, 0) > 1, "DuplicateFactoryNumber")
        add_flag(flags, not pico, "MissingPicoSerial")
        add_flag(flags, pico_counts.get(pico, 0) > 1, "DuplicatePicoSerial")
        add_flag(flags, bool(pico) and len(pico_links.get(pico, ())) > 1, "ConflictingPicoSerial")
        add_flag(flags, not art, "MissingArtSerial")
        add_flag(flags, art_counts.get(art, 0) > 1, "DuplicateArtSerial")
        add_flag(flags, bool(art) and len(art_links.get(art, ())) > 1, "ConflictingArtSerial")
        add_flag(flags, not record["deviceModel"], "MissingDeviceModel")
        add_flag(flags, record["dividerRatio"] is None, "MissingDividerRatio")
        add_flag(
            flags,
            record["inductor"]["inductanceHenry"] is None
            or record["inductor"]["dcResistanceOhm"] is None,
            "MissingInductorParameters",
        )
        add_flag(flags, record["couplingCapacitanceNanoFarad"] is None, "MissingCouplingCapacitance")
        add_flag(flags, record["compensation"]["selected"] is None, "MissingCompensationSelection")
        add_flag(flags, not record["calibrationCoefficientSource"], "MissingCalibrationCoefficient")
        add_flag(
            flags,
            bool(record["calibrationCoefficientSource"]) and record["calibrationCoefficient"] is None,
            "NonNumericCalibrationCoefficient",
        )
        record["validationFlags"] = flags
        record["autoMatchEligible"] = not any(
            flag
            in {
                "MissingFactoryNumber",
                "DuplicateFactoryNumber",
                "MissingPicoSerial",
                "DuplicatePicoSerial",
                "ConflictingPicoSerial",
                "MissingArtSerial",
                "DuplicateArtSerial",
                "ConflictingArtSerial",
            }
            for flag in flags
        )
        if not record["autoMatchEligible"]:
            for flag in flags:
                if flag.startswith(("MissingFactory", "DuplicateFactory", "MissingPico", "DuplicatePico", "ConflictingPico", "MissingArt", "DuplicateArt", "ConflictingArt")):
                    exclusion_counts[flag] += 1

    def identifier_stats(selector):
        values = [selector(r) for r in records]
        values = [value for value in values if value]
        counts = collections.Counter(values)
        return {
            "nonEmptyRows": len(values),
            "distinctKeys": len(counts),
            "uniqueKeys": sum(count == 1 for count in counts.values()),
            "duplicateKeys": sum(count > 1 for count in counts.values()),
            "duplicateRows": sum(count for count in counts.values() if count > 1),
        }

    source_hash = hashlib.sha256(source.read_bytes()).hexdigest().upper()
    generated_date = dt.date.today().isoformat()
    factory_stats = identifier_stats(lambda r: r["factoryNumber"])
    pico_stats = identifier_stats(lambda r: r["picoScope"]["serial"])
    art_stats = identifier_stats(lambda r: r["art"]["serial"])
    pico_stats.update(
        {
            "conflictKeys": sum(len(links) > 1 for links in pico_links.values()),
            "conflictRows": sum(
                pico_counts[key] for key, links in pico_links.items() if len(links) > 1
            ),
        }
    )
    art_stats.update(
        {
            "conflictKeys": sum(len(links) > 1 for links in art_links.values()),
            "conflictRows": sum(
                art_counts[key] for key, links in art_links.items() if len(links) > 1
            ),
        }
    )
    return {
        "schema": "XIONGXIN.DeviceAssetCatalog.v1",
        "catalogVersion": generated_date + ".1",
        "provenance": {
            "sourceDocument": source.name,
            "sourceSha256": source_hash,
            "worksheet": "出厂记录",
            "sourceRange": "A2:Y87",
            "privacyPolicy": "Device and calibration fields only; customer, shipment, repair and free-text note fields omitted.",
        },
        "numberingRules": {
            "version": "1",
            "mainDeviceSerial": {
                "composition": [
                    "switchTypeCode",
                    "pdChannelCode",
                    "syncTriggerCode",
                    "couplingCapacitorCodeAndValue",
                    "dateYYMMDD",
                    "factorySequence",
                    "dcSourceCode",
                ],
                "dateFormat": "YYMMDD",
                "factorySequenceFormat": "# plus three digits",
                "factorySequenceMeaning": "Monthly device sequence matching the factory record number",
            },
            "typeNumber": {
                "composition": ["EPFT-OWTS", "voltageBandPrefix", "voltageLevelWithoutUnit"],
                "voltageBandPrefixes": {
                    "M": "below 60 kV",
                    "H": "above 60 kV (the source does not define exactly 60 kV)",
                },
            },
            "codes": {
                "switchType": {"L": "optical trigger", "M": "magnetic trigger"},
                "pdChannel": {"A": "PD acquisition channel A", "B": "PD acquisition channel B"},
                "syncTrigger": {
                    "A": "channel A synchronous trigger",
                    "B": "channel B synchronous trigger",
                    "E": "digital external trigger",
                    "L": "logic-level trigger X terminal",
                },
                "picoScopeModel": {"P5": "PicoScope 5242D"},
                "artModel": {"Y3": "ART 3135A terminal-block type", "D3": "ART 3135 DB37 type"},
                "couplingCapacitor": {
                    "C": "circuit-board type",
                    "T": "TDK capacitor",
                    "V": "Vishay capacitor",
                },
            },
            "usage": "Documented provenance only; these rules must not infer live hardware identity or overwrite settings.",
        },
        "audit": {
            "sourceRows": len(records),
            "factory": factory_stats,
            "pico": pico_stats,
            "art": art_stats,
            "pairedRows": sum(bool(r["picoScope"]["serial"] and r["art"]["serial"]) for r in records),
            "autoMatchEligibleRows": sum(r["autoMatchEligible"] for r in records),
            "excludedRows": sum(not r["autoMatchEligible"] for r in records),
            "exclusionReasonRows": dict(sorted(exclusion_counts.items())),
        },
        "records": records,
    }


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", help="Path to the source .xlsx workbook (opened read-only)")
    args = parser.parse_args()
    print(json.dumps(generate(args.workbook), ensure_ascii=False, indent=2, sort_keys=False))


if __name__ == "__main__":
    main()
